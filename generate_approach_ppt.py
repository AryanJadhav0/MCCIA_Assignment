import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def main():
    prs = Presentation()

    # -------------------------------------------------------------
    # Helper to apply professional styling to title text
    # -------------------------------------------------------------
    def style_title(title_shape):
        if not title_shape:
            return
        for paragraph in title_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Calibri'
                run.font.bold = True
                run.font.color.rgb = RGBColor(31, 56, 100) # Navy Blue

    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    slide_layout = prs.slide_layouts[0] # Title Slide layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "The Supplier Blindspot"
    subtitle.text = "Analytical Framework, Technical Solution, and Strategic Business Approach\n\nArora Traders Case Study"
    style_title(title)

    # -------------------------------------------------------------
    # SLIDE 2: THE BUSINESS PROBLEM
    # -------------------------------------------------------------
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "1. The Business Challenge"
    style_title(title)
    
    body = slide.shapes.placeholders[1].text_frame
    body.text = "The Objective:"
    
    p = body.add_paragraph()
    p.text = "Identify the root cause of financial leakages and supply chain inefficiencies among 34 suppliers over a 3-year period (4,800 POs)."
    p.level = 1
    
    p = body.add_paragraph()
    p.text = "The 'Blindspot':"
    p.level = 0
    
    p = body.add_paragraph()
    p.text = "Poor data visibility, off-master POs, and fragmented operations masked which suppliers were actually costing the business the most money."
    p.level = 1

    # -------------------------------------------------------------
    # SLIDE 3: OUR CORE APPROACH
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "2. The Multi-Stage Data Pipeline"
    style_title(title)
    
    body = slide.shapes.placeholders[1].text_frame
    body.text = "How we transformed raw data into actionable intelligence:"
    
    p = body.add_paragraph()
    p.text = "Step 1: Attribution & Unification"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Cleaned and mapped messy, off-master POs to their correct supplier IDs using algorithmic matching."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "Step 2: Financial Impact Modeling"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Transitioned away from vague operational metrics to hard financial losses."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "Step 3: Composite Scoring Engine"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Weighted scoring algorithm evaluating delivery, quality, and pricing."
    p.level = 2

    # -------------------------------------------------------------
    # SLIDE 4: FINANCIAL MODELING
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "3. Financial Modeling (The A+B+C+D Framework)"
    style_title(title)
    
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Every operational failure was quantified into a unified Headline Rupee Loss:"
    
    p = body.add_paragraph()
    p.text = "A*: Excess Billing Gap"
    p.level = 1
    p = body.add_paragraph()
    p.text = "When invoice quantities exceeded received quantities beyond an acceptable baseline."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "B: Rejection Losses"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Material billed in full but subsequently rejected."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "C: Confirmed Returns"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Directly traced returned batches mapped to specific POs."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "D: Inferred Returns"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Untraced returns probabilistically allocated."
    p.level = 2

    # -------------------------------------------------------------
    # SLIDE 5: OUT-OF-THE-BOX PROBLEM SOLVING
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "4. Out-of-the-Box Business Strategy"
    style_title(title)
    
    body = slide.shapes.placeholders[1].text_frame
    body.text = "Strategic differentiation in our approach:"
    
    p = body.add_paragraph()
    p.text = "Probabilistic Financial Attribution"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Instead of writing off untraced returns, we assigned financial penalties based on historical quality failure rates to ensure accountability."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "Unified Rupee Impact"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Instead of reporting 'Supplier X is 5 days late', we converted every failure into an absolute financial loss to force immediate executive prioritization."
    p.level = 2

    # -------------------------------------------------------------
    # SLIDE 6: BRIDGING DATA & EXECUTION
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "5. Bridging Data & Execution"
    style_title(title)
    
    body = slide.shapes.placeholders[1].text_frame
    body.text = "We didn't just stop at a dashboard of numbers."
    
    p = body.add_paragraph()
    p.text = "Automated Negotiation Briefs"
    p.level = 1
    
    p = body.add_paragraph()
    p.text = "We auto-generated structured strategy documents for the poorest performing suppliers."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "Empowers procurement managers with exact negotiation scripts."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "Creates corrective Action Plans tailored to the specific weakness of the supplier (e.g., if pricing is the issue, it pushes for panel-matching)."
    p.level = 2

    # -------------------------------------------------------------
    # SLIDE 7: THE FINAL SOLUTION
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "6. The Delivered Solution"
    style_title(title)
    
    body = slide.shapes.placeholders[1].text_frame
    body.text = "The culmination of our analysis:"
    
    p = body.add_paragraph()
    p.text = "Zero-Backend Interactive Dashboard"
    p.level = 1
    p = body.add_paragraph()
    p.text = "A lightning-fast, 100% client-side web application deployed to the cloud."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "34-Supplier Scorecard"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Searchable categorization of suppliers into POOR, WATCH, and ACCEPTABLE tiers."
    p.level = 2
    
    p = body.add_paragraph()
    p.text = "Visualized Financial Impact"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Dynamic charts exposing exact loss compositions to drive immediate ROI."
    p.level = 2

    # -------------------------------------------------------------
    # SLIDE 8: THANK YOU PAGE
    # -------------------------------------------------------------
    slide_layout = prs.slide_layouts[0] # Title Slide layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Thank You"
    subtitle.text = "Questions & Discussion"
    style_title(title)

    # Save presentation
    prs.save("Arora_Traders_Approach_Presentation.pptx")
    print("Successfully generated Arora_Traders_Approach_Presentation.pptx")

if __name__ == '__main__':
    main()
